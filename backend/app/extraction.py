"""Read student documents into ordered, reviewable source units.

When configured, the document AI provider reads image pages and normalizes text.
Tesseract remains the local fallback. The original upload is retained by the
job queue; only the configured external provider receives document content.
"""

import io
import logging
import os
import re
import subprocess
import tempfile
import unicodedata
import zipfile
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 150_000
MAX_PDF_PAGES = 30
MAX_IMAGE_PIXELS = 20_000_000
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}
WORD_SUFFIXES = {".doc", ".docx"}
TEXT_DOCUMENT_SUFFIXES = {".txt", ".md", ".csv", ".rtf", ".html", ".htm"}
logger = logging.getLogger(__name__)


class ExtractionError(ValueError):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


def validate_upload(filename: str, data: bytes):
    suffix = Path(filename).suffix.lower()
    if suffix not in {*TEXT_DOCUMENT_SUFFIXES, ".pdf", ".doc", ".docx", ".odt", *IMAGE_SUFFIXES}:
        raise ExtractionError("Chỉ hỗ trợ TXT, PDF, DOCX, PNG, JPG/JPEG và WEBP.", 415)
    if not data:
        raise ExtractionError("Tệp rỗng")
    if len(data) > MAX_FILE_BYTES:
        raise ExtractionError("Tệp vượt quá 10MB", 413)
    if suffix == ".pdf" and not data.startswith(b"%PDF-"):
        raise ExtractionError("Tệp PDF không hợp lệ")
    if suffix == ".docx" and not zipfile.is_zipfile(io.BytesIO(data)):
        raise ExtractionError("Tệp DOCX không hợp lệ", 415)
    if suffix == ".doc" and data[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ExtractionError("Tệp DOC không hợp lệ", 415)
    if suffix == ".odt" and not zipfile.is_zipfile(io.BytesIO(data)):
        raise ExtractionError("Tệp ODT không hợp lệ", 415)
    if suffix == ".png" and not data.startswith(b"\x89PNG\r\n\x1a\n"):
        raise ExtractionError("Tệp PNG không hợp lệ", 415)
    if suffix in {".jpg", ".jpeg"} and not data.startswith(b"\xff\xd8"):
        raise ExtractionError("Tệp JPG không hợp lệ", 415)
    if suffix == ".webp" and not (data.startswith(b"RIFF") and data[8:12] == b"WEBP"):
        raise ExtractionError("Tệp WEBP không hợp lệ", 415)


def _image(data: bytes):
    from PIL import Image, ImageOps, UnidentifiedImageError

    try:
        image = Image.open(io.BytesIO(data))
        if image.width * image.height > MAX_IMAGE_PIXELS:
            raise ExtractionError("Ảnh quá lớn để OCR; tối đa 20 triệu điểm ảnh")
        image = ImageOps.exif_transpose(image)
        image.load()
        return image.convert("RGB")
    except ExtractionError:
        raise
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise ExtractionError("Không đọc được ảnh; hãy kiểm tra định dạng hoặc chụp lại") from exc


def _prepare_ocr_image(image):
    """Create a stable OCR view without changing the stored upload.

    The order is deliberate: remove colour, upscale small scans, suppress
    speckle noise, normalize the page range, then apply a restrained sharpen.
    We keep grayscale instead of hard thresholding so Vietnamese tone marks
    are not erased.
    """
    from PIL import ImageEnhance, ImageFilter, ImageOps

    prepared = ImageOps.grayscale(image)
    if prepared.width < 1800:
        scale = min(1.6, 1800 / max(prepared.width, 1))
        prepared = prepared.resize((round(prepared.width * scale), round(prepared.height * scale)))
    prepared = prepared.filter(ImageFilter.MedianFilter(size=3))
    prepared = ImageOps.autocontrast(prepared, cutoff=1)
    prepared = ImageEnhance.Contrast(prepared).enhance(1.25)
    return prepared.filter(ImageFilter.UnsharpMask(radius=1.2, percent=140, threshold=3))


def _ocr_image(image) -> tuple[str, float | None]:
    import pytesseract
    from PIL import ImageEnhance, ImageOps
    from pytesseract import Output

    if os.getenv("OCR_TESSERACT_CMD"):
        pytesseract.pytesseract.tesseract_cmd = os.environ["OCR_TESSERACT_CMD"]

    prepared = _prepare_ocr_image(image)

    try:
        words = pytesseract.image_to_data(
            prepared, lang="vie", config="--psm 3", output_type=Output.DICT, timeout=35,
        )
    except pytesseract.TesseractNotFoundError as exc:
        raise ExtractionError("Máy chủ chưa cài Tesseract OCR; cần cài Tesseract và dữ liệu tiếng Việt 'vie'.") from exc
    except (pytesseract.TesseractError, RuntimeError) as exc:
        raise ExtractionError("OCR ảnh thất bại hoặc quá thời gian; hãy thử ảnh rõ hơn.") from exc

    def mean_confidence(candidate):
        values = []
        for value in candidate.get("conf", []):
            try:
                score = float(value)
            except (TypeError, ValueError):
                continue
            if score >= 0:
                values.append(score)
        return sum(values) / len(values) if values else 0

    # A second layout pass and per-word proofreading are expensive on scanned
    # PDFs. Keep the default bounded so one bad page cannot hold the worker for
    # minutes; opt into the slower proofreading pass explicitly when auditing
    # a difficult scan.
    detailed = os.getenv("OCR_DETAILED_VERIFY", "0").strip().lower() in {"1", "true", "yes"}
    first_score = mean_confidence(words)
    if detailed and first_score < 88:
        try:
            alternate = pytesseract.image_to_data(
                prepared, lang="vie", config="--psm 6", output_type=Output.DICT, timeout=35,
            )
            if mean_confidence(alternate) > first_score + 2:
                words = alternate
        except (pytesseract.TesseractError, RuntimeError):
            pass

    lines = []
    current_key = None
    line_words = []
    line_indexes = []
    scores = []

    def plain(value: str) -> str:
        decomposed = unicodedata.normalize("NFD", value.upper().replace("Đ", "D"))
        return "".join(char for char in decomposed if unicodedata.category(char) != "Mn")

    def accent_count(value: str) -> int:
        return sum(unicodedata.category(char) == "Mn" for char in unicodedata.normalize("NFD", value))

    def verified_number(index: int, value: str, confidence: float) -> str:
        """Check uncertain numerals on their own pixels with digit-only OCR."""
        if not detailed:
            return value
        numeric = re.fullmatch(r"(\d{1,6})([.,):%]*)", value)
        after_question_label = bool(line_words and plain(line_words[-1]).rstrip(".") in {"CAU", "BAI"})
        ambiguous_label_number = (after_question_label and bool(re.fullmatch(r"\D{1,3}[.):]", value)))
        if numeric is None and not ambiguous_label_number:
            return value
        if numeric is not None and confidence >= 99:
            return value
        left, top = words["left"][index], words["top"][index]
        right = left + words["width"][index]
        bottom = top + words["height"][index]
        crop = prepared.crop((max(0, left - 8), max(0, top - 8),
                              min(prepared.width, right + 8), min(prepared.height, bottom + 8)))
        try:
            reads = [pytesseract.image_to_string(
                crop, lang="eng", config=f"--psm {psm} -c tessedit_char_whitelist=0123456789",
                timeout=5).strip() for psm in (7, 8)]
        except (pytesseract.TesseractError, RuntimeError):
            return value
        if reads[0] != reads[1] or not reads[0].isdigit():
            return value
        if numeric is not None:
            if len(reads[0]) != len(numeric.group(1)):
                return value
            return reads[0] + numeric.group(2)
        return reads[0] + value[len(value.rstrip(".,):%")):]

    def verified_word(index: int, value: str, confidence: float) -> str:
        """Recheck low-confidence Vietnamese words without guessing from context."""
        if not detailed:
            return value
        if confidence >= 70 or not re.fullmatch(r"[^\W\d_]{2,20}[.!?,;:]?", value, re.UNICODE):
            return value
        left, top = words["left"][index], words["top"][index]
        right = left + words["width"][index]
        bottom = top + words["height"][index]
        crop = prepared.crop((max(0, left - 12), max(0, top - 8),
                              min(prepared.width, right + 12), min(prepared.height, bottom + 8)))
        try:
            reads = [pytesseract.image_to_string(
                crop, lang="vie", config=f"--psm {psm}", timeout=5,
            ).strip().strip("\"'“”‘’.,!?;:") for psm in (7, 8)]
        except (pytesseract.TesseractError, RuntimeError):
            return value
        original = value.rstrip(".!?,;:")
        if reads[0] != reads[1] or not re.fullmatch(r"[^\W\d_]{2,20}", reads[0], re.UNICODE):
            return value
        if len(reads[0]) < len(original) - 1 or len(reads[0]) > len(original) + 1:
            return value
        return reads[0] + value[len(original):]

    def flush():
        if line_words:
            value = " ".join(line_words)
            # Short all-caps headings are especially prone to missed Vietnamese
            # accents. Retry only their own crop, never replace unrelated lines.
            if (detailed and value.isupper() and 2 <= len(line_words) <= 12 and len(value) <= 80
                    and not any(char.isdigit() for char in value)):
                left = min(words["left"][i] for i in line_indexes)
                top = min(words["top"][i] for i in line_indexes)
                right = max(words["left"][i] + words["width"][i] for i in line_indexes)
                bottom = max(words["top"][i] + words["height"][i] for i in line_indexes)
                crop = prepared.crop((max(0, left - 30), max(0, top - 30),
                                      min(prepared.width, right + 30), min(prepared.height, bottom + 30)))
                prepared_crop = ImageEnhance.Contrast(ImageOps.grayscale(crop)).enhance(1.6)
                try:
                    candidate = " ".join(pytesseract.image_to_string(
                        prepared_crop, lang="vie", config="--psm 6", timeout=10).split())
                    if plain(candidate) == plain(value) and accent_count(candidate) > accent_count(value):
                        value = candidate
                except (pytesseract.TesseractError, RuntimeError):
                    pass
            lines.append(value)
            line_words.clear()
            line_indexes.clear()

    for index, value in enumerate(words["text"]):
        value = value.strip()
        if not value:
            continue
        try:
            confidence = float(words["conf"][index])
        except (ValueError, TypeError):
            confidence = -1
        if confidence < 0:
            continue
        key = (words["block_num"][index], words["par_num"][index], words["line_num"][index])
        if current_key is not None and key != current_key:
            flush()
        value = verified_number(index, value, confidence)
        value = verified_word(index, value, confidence)
        current_key = key
        line_words.append(value)
        line_indexes.append(index)
        scores.append(confidence)
    flush()
    return "\n".join(lines).strip(), round(sum(scores) / len(scores), 1) if scores else None


def _jpeg_bytes(image) -> bytes:
    # Keep model payloads bounded. The original upload remains untouched; this
    # representation is only used for OCR/vision requests and previews.
    image = image.convert("RGB")
    max_dimension = 1800
    if max(image.size) > max_dimension:
        scale = max_dimension / max(image.size)
        image = image.resize((round(image.width * scale), round(image.height * scale)))
    output = io.BytesIO()
    image.save(output, format="JPEG", quality=72, optimize=True)
    return output.getvalue()


def _pdf_pages(data: bytes, source_images: dict[int, bytes] | None = None,
               visual_sources: list[dict] | None = None, analyze_visuals: bool = False):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted or len(reader.pages) > MAX_PDF_PAGES:
        raise ExtractionError("PDF phải không mã hóa và tối đa 30 trang")
    rendered = None
    pages = []
    try:
        for index, page in enumerate(reader.pages):
            # With a vision provider configured, classify and route PDF pages
            # from their rendered images before extracting any page text.
            text = "" if source_images is not None else (page.extract_text() or "").strip()
            method = "text"
            confidence = None
            needs_ocr = len(text) < 15 and (not text or len(page.images) > 0)
            # Text-native PDFs should stay text-native. Rendering every page
            # just because visual analysis is enabled makes a normal PDF look
            # like an image upload and creates unnecessary vision requests.
            # Render only scans, embedded-image pages, or very short text pages
            # that may contain a diagram/formula without a text layer.
            visual_candidate = analyze_visuals and (bool(page.images) or len(text) < 80)
            classify_page = source_images is not None
            if needs_ocr or visual_candidate or classify_page:
                import pypdfium2 as pdfium

                if rendered is None:
                    rendered = pdfium.PdfDocument(data)
                image = rendered[index].render(scale=1.5).to_pil()
                if image.width * image.height > MAX_IMAGE_PIXELS:
                    raise ExtractionError(f"Trang {index + 1} quá lớn để OCR")
            if visual_sources is not None and (needs_ocr or visual_candidate or classify_page):
                visual_sources.append({"page": index + 1, "source_id": f"page-{index + 1}",
                                       "source": "pdf_page",
                                       "image": _jpeg_bytes(image)})
            if classify_page:
                source_images[index + 1] = _jpeg_bytes(image)
            if needs_ocr:
                if source_images is not None:
                    text, confidence = "", None
                else:
                    text, confidence = _ocr_image(image)
                method = "ocr"
            pages.append({"number": index + 1, "label": f"Trang {index + 1}", "text": text,
                          "method": method, "confidence": confidence,
                          **({"image_source": "pdf_page"} if classify_page else {})})
    finally:
        if rendered is not None:
            rendered.close()
    return pages


def _docx_pages(data: bytes):
    from docx import Document

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
            raise ExtractionError("DOCX có kích thước giải nén quá lớn")
    doc = Document(io.BytesIO(data))
    blocks = []
    for block in doc.iter_inner_content():
        if hasattr(block, "rows"):
            blocks.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
        else:
            blocks.append(block.text)
    return [{"number": 1, "label": "Tài liệu Word", "text": "\n\n".join(blocks).strip(),
             "method": "text", "confidence": None}]


def _office_page_images(data: bytes, suffix: str):
    """Render Office documents into ordered page images for layout analysis."""
    from pypdf import PdfReader
    import pypdfium2 as pdfium

    office = os.getenv("LIBREOFFICE_BIN", "soffice")
    with tempfile.TemporaryDirectory(prefix="socratic-pages-") as directory:
        source = Path(directory) / f"input{suffix}"
        source.write_bytes(data)
        try:
            completed = subprocess.run(
                [office, "--headless", "--convert-to", "pdf", "--outdir", directory, str(source)],
                capture_output=True, text=True, timeout=60, check=False,
            )
        except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
            raise ExtractionError("LibreOffice is required to render this file for layout analysis.", 415) from exc
        rendered_path = Path(directory) / "input.pdf"
        if completed.returncode != 0 or not rendered_path.exists():
            raise ExtractionError("Could not render the Office file for layout analysis.", 415)
        reader = PdfReader(str(rendered_path))
        if reader.is_encrypted or not 1 <= len(reader.pages) <= MAX_PDF_PAGES:
            raise ExtractionError(f"Rendered document must contain 1 to {MAX_PDF_PAGES} pages.")
        document = pdfium.PdfDocument(str(rendered_path))
        page_images, page_records, sources = {}, [], []
        try:
            for index in range(len(document)):
                image = document[index].render(scale=1.5).to_pil()
                if image.width * image.height > MAX_IMAGE_PIXELS:
                    raise ExtractionError(f"Rendered page {index + 1} is too large for layout analysis.")
                number = index + 1
                image_bytes = _jpeg_bytes(image)
                page_images[number] = image_bytes
                page_records.append({"number": number, "label": f"Page {number}",
                                     "text": reader.pages[index].extract_text() or "",
                                     "method": "text", "confidence": None,
                                     "image_source": "office_page"})
                sources.append({"page": number, "source_id": f"page-{number}",
                                "source": "office_page", "image": image_bytes})
        finally:
            document.close()
        return page_records, page_images, sources


def _docx_visual_sources(data: bytes) -> list[dict]:
    """Keep embedded media addressable even though DOCX has no stable page numbers."""
    sources = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for entry in archive.infolist():
            if not entry.filename.startswith("word/media/") or entry.is_dir():
                continue
            if entry.file_size > MAX_FILE_BYTES or len(sources) >= 30:
                continue
            try:
                image = _image(archive.read(entry))
            except ExtractionError:
                continue
            sources.append({"page": 1, "source_id": f"media-{len(sources) + 1}",
                            "source": "docx_media", "media_path": entry.filename,
                            "image": _jpeg_bytes(image)})
    return sources


def _doc_pages(data: bytes):
    """Convert legacy binary DOC through LibreOffice, then parse the DOCX."""
    office = os.getenv("LIBREOFFICE_BIN", "soffice")
    with tempfile.TemporaryDirectory(prefix="socratic-doc-") as directory:
        source = Path(directory) / "input.doc"
        source.write_bytes(data)
        try:
            completed = subprocess.run(
                [office, "--headless", "--convert-to", "docx", "--outdir", directory, str(source)],
                capture_output=True, text=True, timeout=45, check=False,
            )
        except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
            raise ExtractionError("Chưa cài LibreOffice để đổi DOC sang DOCX; hãy tải DOCX hoặc PDF.", 415) from exc
        converted = Path(directory) / "input.docx"
        if completed.returncode != 0 or not converted.exists():
            raise ExtractionError("Không chuyển được DOC cũ; hãy lưu lại dưới dạng DOCX hoặc PDF.", 415)
        return _docx_pages(converted.read_bytes())


def _decode_text(data: bytes) -> str:
    """Decode common Vietnamese text encodings without silently replacing bytes."""
    for encoding in ("utf-8-sig", "utf-16", "cp1258", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Không đọc được mã hóa TXT; hãy lưu tài liệu dưới dạng UTF-8.", 422)


class _HTMLTextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript"}:
            self.skip += 1
        elif tag in {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr"} and self.parts:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"} and self.skip:
            self.skip -= 1
        elif tag in {"p", "div", "li", "h1", "h2", "h3", "h4", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def _odt_text(data: bytes) -> str:
    import xml.etree.ElementTree as ET

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            content = archive.read("content.xml")
        root = ET.fromstring(content)
    except (KeyError, OSError, ET.ParseError) as exc:
        raise ExtractionError("Không đọc được nội dung ODT.", 422) from exc
    values = []
    for node in root.iter():
        tag = node.tag.rsplit("}", 1)[-1]
        if tag in {"p", "h", "list-item", "table-cell"}:
            value = "".join(node.itertext()).strip()
            if value:
                values.append(value)
    return "\n\n".join(values)


def _rich_text(data: bytes, suffix: str) -> str:
    text = _decode_text(data)
    if suffix in {".html", ".htm"}:
        parser = _HTMLTextParser()
        parser.feed(text)
        text = "".join(parser.parts)
    elif suffix == ".rtf":
        # Keep readable text while removing common RTF control words and groups.
        text = re.sub(r"\\'[0-9a-fA-F]{2}", "", text)
        text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
        text = text.replace("{", "").replace("}", "")
    return text


def classify_document(text: str, filename: str = "") -> str:
    """Coarse profile used to select a subject/document parser."""
    folded = unicodedata.normalize("NFD", text.upper())
    folded = "".join(c for c in folded if unicodedata.category(c) != "Mn")
    if any(marker in folded for marker in ("DAP AN", "HUONG DAN CHAM", "THANG DIEM", "RUBRIC")):
        return "answer_key"
    if any(marker in folded for marker in ("DE THI", "CAU 1", "CAU 2", "THOI GIAN LAM BAI")):
        return "exam"
    if any(marker in folded for marker in ("MUC TIEU BAI HOC", "HOAT DONG", "LUYEN TAP")):
        return "lesson_note"
    if any(marker in folded for marker in ("BAI TAP", "YEU CAU", "NHIEM VU")):
        return "assignment"
    if filename.lower().endswith((".jpg", ".jpeg", ".png", ".webp")):
        return "image_note"
    return "reference_document"


def detect_subject(text: str, filename: str = "") -> str:
    """Detect a conservative subject hint; it never replaces user confirmation."""
    folded = unicodedata.normalize("NFD", (text + " " + filename).upper())
    folded = "".join(c for c in folded if unicodedata.category(c) != "Mn")
    markers = {
        "ngữ văn": ("NGU VAN", "DOC HIEU", "NGHI LUAN", "BAI THO", "THO"),
        "toán": ("TOAN", "PHUONG TRINH", "HAM SO", "HINH HOC", "TICH PHAN"),
        "vật lý": ("VAT LY", "CO HOC", "DIEN XOAY CHIEU", "QUANG HOC"),
        "hóa học": ("HOA HOC", "PHAN UNG HOA HOC", "OXI HOA", "ESTE"),
        "sinh học": ("SINH HOC", "DI TRUYEN", "TE BAO", "SINH THAI"),
        "lịch sử": ("LICH SU", "SU KIEN", "KHOI NGHIA", "CHIEN TRANH"),
        "địa lý": ("DIA LY", "KHI HAU", "DAN SO", "DIA HINH"),
        "tiếng Anh": ("ENGLISH", "READING", "GRAMMAR", "LISTENING", "VOCABULARY"),
        "tin học": ("TIN HOC", "LAP TRINH", "THUAT TOAN", "PYTHON", "HTML"),
    }
    for subject, candidates in markers.items():
        if any(candidate in folded for candidate in candidates):
            return subject
    return "chưa xác định"


def subject_from_pages(pages: list[dict], text: str, filename: str = "") -> str:
    """Prefer consistent model hints, retaining deterministic detection as fallback."""
    subjects = [page["subject_hint"].strip() for page in pages
                if isinstance(page.get("subject_hint"), str) and page["subject_hint"].strip()]
    counts = Counter(subjects).most_common(2)
    return counts[0][0] if counts and (len(counts) == 1 or counts[0][1] > counts[1][1]) \
        else detect_subject(text, filename)


def _content_blocks(pages: list[dict]) -> list[dict]:
    """Create small, provenance-aware blocks for non-exam documents too."""
    blocks = []
    for page in pages:
        paragraph = []

        def flush():
            if paragraph:
                value = " ".join(part.strip() for part in paragraph).strip()
                if value:
                    blocks.append({"id": f"block-{len(blocks) + 1}", "type": "paragraph",
                                   "text": value, "page": page["number"], "method": page["method"]})
                paragraph.clear()

        for raw in page["text"].replace("\r", "").split("\n"):
            line = raw.strip()
            if not line:
                flush()
                continue
            if len(line) <= 100 and (line.isupper() or line.endswith(":") or line.startswith(("#", "##"))):
                flush()
                blocks.append({"id": f"block-{len(blocks) + 1}", "type": "heading",
                               "text": line.lstrip("# "), "page": page["number"], "method": page["method"]})
            elif line.startswith(("- ", "• ", "* ")):
                flush()
                blocks.append({"id": f"block-{len(blocks) + 1}", "type": "list_item",
                               "text": line[2:].strip(), "page": page["number"], "method": page["method"]})
            else:
                paragraph.append(line)
        flush()
    return blocks


def _finalize_pages(filename: str, pages: list[dict], source_images: dict[int, bytes] | None) -> dict:
    from .document_ai import enabled as document_ai_enabled, improve_pages
    from concurrent.futures import ThreadPoolExecutor, as_completed

    ai_warnings = []
    natural_pages = set()
    unclassified_image_pages = set()
    if source_images and document_ai_enabled():
        from .visual_layout import analyze_question_document
        candidates = [page for page in pages if page["number"] in source_images]
        try:
            analysis = analyze_question_document([
                {"page": page["number"], "source_id": f"page-{page['number']}",
                 "source": page.get("image_source", "image_page"),
                 "image": source_images[page["number"]]}
                for page in candidates
            ])
        except Exception as exc:
            logger.warning("Question document analysis failed pages=%d error=%s",
                           len(candidates), type(exc).__name__)
            unclassified_image_pages.update(page["number"] for page in candidates)
            ai_warnings.append("Could not classify document pages from images; keeping the existing text flow.")
        else:
            unresolved = analysis.get("unresolved_regions", [])
            unresolved_pages = sorted({item.get("page") for item in unresolved
                                       if isinstance(item, dict) and item.get("page") in analysis["pages"]})
            if unresolved_pages:
                ai_warnings.append("Some image regions could not be assigned confidently on pages: "
                                   + ", ".join(map(str, unresolved_pages)) + ".")
            for page in candidates:
                page_analysis = analysis["pages"].get(page["number"], {})
                page["subject_hint"] = analysis["subject_hint"]
                if analysis["subject_group"] != "natural":
                    continue
                natural_pages.add(page["number"])
                page["method"] = "image"
                page["text"] = ""
                page["visuals"] = page_analysis.get("questions", [])
                if analysis["subject_uncertain"]:
                    ai_warnings.append(f"{page['label']}: subject classification is uncertain; retaining question content as images.")
                if not page["visuals"]:
                    ai_warnings.append(f"{page['label']}: no question frames detected; review the source image.")
                elif any(question["uncertain"] for question in page["visuals"]):
                    ai_warnings.append(f"{page['label']}: some question frames or page continuations are uncertain; review the source image.")

    text_pages = [page for page in pages if page["number"] not in natural_pages
                  and page["number"] not in unclassified_image_pages]
    text_images = {number: image for number, image in (source_images or {}).items()
                   if number not in natural_pages}
    pages_by_number = {page["number"]: page for page in pages}
    if text_pages:
        improved, warnings = improve_pages(text_pages, text_images)
        ai_warnings.extend(warnings)
        for page in improved:
            pages_by_number[page["number"]].update(page)
    if source_images:
        pending = [page for page in pages if not page["text"] and page["number"] in source_images
                   and page["number"] not in natural_pages]
        # OCR pages independently; Tesseract runs in child processes, so a
        # small pool cuts a multi-page scan from serial page latency to roughly
        # two or three waves while preserving the original order in `pages`.
        with ThreadPoolExecutor(max_workers=min(3, max(1, len(pending)))) as pool:
            futures = {pool.submit(_ocr_image, _image(source_images[page["number"]])): page
                       for page in pending}
            for future in as_completed(futures):
                page = futures[future]
                try:
                    page["text"], page["confidence"] = future.result()
                    page["method"] = "ocr"
                    ai_warnings.append(f"{page['label']}: đã dùng OCR cục bộ vì AI không đọc được trang này.")
                except ExtractionError:
                    pass
    text = "\n\n".join(page["text"] for page in pages if page["text"]).strip()
    has_image_questions = any(page.get("visuals") and page["number"] in natural_pages for page in pages)
    if len(text) < 15 and not has_image_questions:
        if natural_pages:
            raise ExtractionError("AI chưa khoanh được câu hỏi trong ảnh; hãy thử ảnh rõ hơn hoặc chia ảnh theo từng câu.")
        raise ExtractionError("Không nhận diện đủ chữ; hãy dùng ảnh rõ hơn hoặc nhập nội dung trực tiếp.")
    if len(text) > MAX_TEXT_CHARS:
        raise ExtractionError("Nội dung vượt quá 150.000 ký tự; hãy chia tài liệu thành nhiều bộ đề.")
    warnings = [*ai_warnings, *[f"{page['label']} không đọc được chữ; cần kiểm tra bản gốc."
                                 for page in pages if not page["text"] and page["method"] != "image"]]
    if any(page["method"] == "ocr" for page in pages):
        warnings.append("OCR có thể sai dù điểm tin cậy cao; hãy đối chiếu các số liệu và chữ quan trọng với bản gốc.")
    warnings.extend(f"{page['label']} OCR chưa chắc chắn; cần xem lại chữ."
                    for page in pages if page["confidence"] is not None and page["confidence"] < 65)
    warnings = list(dict.fromkeys(warnings))
    return {"text": text, "pages": pages, "blocks": _content_blocks(pages),
            "warnings": warnings, "document_type": classify_document(text, filename),
            "subject": subject_from_pages(pages, text, filename)}


def _with_visuals(result: dict, sources: list[dict]) -> dict:
    """Attach image provenance and optional AI regions without losing extracted text."""
    from concurrent.futures import ThreadPoolExecutor, as_completed
    import time
    from .config import positive_integer
    from .visual_layout import analyze, enabled

    pages = {page["number"]: page for page in result["pages"]}
    use_model = enabled()
    limit = positive_integer("DOCUMENT_VISUAL_MAX_SOURCES", 30)
    timeout = positive_integer("DOCUMENT_AI_TIMEOUT_SECONDS", 30)
    deadline = time.monotonic() + max(0, positive_integer("WORKER_JOB_TIMEOUT_SECONDS", 300) - 10)
    failed = 0
    pending = []
    for index, source in enumerate(sources):
        page = pages[source["page"]]
        descriptor = {key: source[key] for key in ("source_id", "source", "media_path") if key in source}
        page.setdefault("visual_sources", []).append(descriptor)
        page.setdefault("visuals", [])
        if any(item.get("kind") == "question" for item in page["visuals"]):
            continue
        if not use_model or not source.get("image"):
            if source["source"] == "docx_media":
                page["visuals"].append({"id": f"{source['source_id']}-region-1",
                    "page": page["number"], "source_id": source["source_id"],
                    "source": "docx_media", "kind": "other", "bbox": [0, 0, 1000, 1000],
                    "orientation": "unknown", "question_label": None,
                    "description": "Hình nhúng trong tài liệu Word", "uncertain": True,
                    "link_status": "unreviewed"})
            continue
        # A scanned page is technically an image, but most pages contain only
        # ordinary text. Do not spend another vision request on those pages;
        # keep the source image and analyze on-demand when the OCR/context
        # suggests a diagram, graph, table, formula, or figure.
        page_text = page.get("text", "")
        visual_hint = re.search(
            r"(hình|biểu đồ|đồ thị|sơ đồ|bảng|công thức|hình vẽ|cho hình|dựa vào hình|quan sát hình)",
            page_text,
            flags=re.IGNORECASE,
        )
        if len(page_text) >= 80 and not visual_hint and source.get("source") in {"pdf_page", "office_page"}:
            continue
        if index >= limit or time.monotonic() + timeout > deadline:
            failed += 1
            continue
        pending.append((page, source))
    if pending:
        # Vision calls are independent. A small pool reduces wall-clock time while
        # staying below typical provider rate limits and preserving page order.
        with ThreadPoolExecutor(max_workers=min(3, len(pending))) as pool:
            futures = {
                pool.submit(analyze, source["image"], page["text"], page["number"],
                            source["source_id"], source["source"]): (page, source)
                for page, source in pending
            }
            results = []
            for future in as_completed(futures):
                page, source = futures[future]
                try:
                    results.append((page["number"], future.result()))
                except Exception:
                    # The external model is advisory; a failure must not lose the upload.
                    failed += 1
            for page_number, visuals in sorted(results, key=lambda item: item[0]):
                page = pages[page_number]
                page["visuals"].extend(visuals)
    if failed:
        result["warnings"].append("Một số hình chưa được phân tích; bản gốc vẫn có thể xem để kiểm tra.")
    return result


def extract_image_batch(files: list[tuple[str, bytes]]) -> dict:
    """Read an ordered set of uploaded images as one document."""
    if not files or len(files) > 12:
        raise ExtractionError("Cần từ 1 đến 12 ảnh trong một tài liệu.")
    from .document_ai import enabled as document_ai_enabled

    source_images = {} if document_ai_enabled() else None
    pages = []
    visual_sources = []
    for number, (filename, data) in enumerate(files, start=1):
        validate_upload(filename, data)
        if Path(filename).suffix.lower() not in IMAGE_SUFFIXES:
            raise ExtractionError("Bộ ảnh chỉ nhận PNG, JPG/JPEG hoặc WEBP.", 415)
        image = _image(data)
        visual_sources.append({"page": number, "source_id": f"page-{number}",
                               "source": "image_page", "image": _jpeg_bytes(image)})
        if source_images is not None:
            source_images[number] = _jpeg_bytes(image)
            text, confidence = "", None
        else:
            text, confidence = _ocr_image(image)
        pages.append({"number": number, "label": f"Ảnh {number}: {filename}", "text": text,
                      "method": "ocr", "confidence": confidence, "image_source": "image_page"})
    return _with_visuals(_finalize_pages("bộ ảnh", pages, source_images), visual_sources)


def extract_document(filename: str, data: bytes) -> dict:
    """Return text and ordered units; never silently discard an unreadable page."""
    validate_upload(filename, data)
    suffix = Path(filename).suffix.lower()
    from .document_ai import enabled as document_ai_enabled
    from .visual_layout import enabled as visual_analysis_enabled
    source_images = {} if document_ai_enabled() else None
    visual_sources = []
    try:
        if suffix in TEXT_DOCUMENT_SUFFIXES:
            text = _rich_text(data, suffix)
            pages = [{"number": 1, "label": f"Tệp {suffix[1:].upper()}", "text": text.strip(),
                      "method": "text", "confidence": None}]
        elif suffix == ".pdf":
            pages = _pdf_pages(data, source_images, visual_sources, visual_analysis_enabled())
        elif suffix == ".docx":
            pages = _docx_pages(data)
            visual_sources = _docx_visual_sources(data)
            if source_images is not None:
                try:
                    pages, rendered_images, rendered_sources = _office_page_images(data, suffix)
                except ExtractionError as exc:
                    logger.warning("Could not render DOCX pages for layout analysis: %s", exc)
                else:
                    source_images.update(rendered_images)
                    visual_sources = [*rendered_sources, *visual_sources]
        elif suffix == ".doc":
            if source_images is not None:
                try:
                    pages, rendered_images, visual_sources = _office_page_images(data, suffix)
                except ExtractionError as exc:
                    logger.warning("Could not render DOC pages for layout analysis: %s", exc)
                    pages = _doc_pages(data)
                else:
                    source_images.update(rendered_images)
            else:
                pages = _doc_pages(data)
        elif suffix == ".odt":
            text = _odt_text(data)
            pages = [{"number": 1, "label": f"File {suffix[1:].upper()}", "text": text.strip(),
                      "method": "text", "confidence": None}]
            if source_images is not None:
                try:
                    pages, rendered_images, visual_sources = _office_page_images(data, suffix)
                except ExtractionError as exc:
                    logger.warning("Could not render ODT pages for layout analysis: %s", exc)
                else:
                    source_images.update(rendered_images)
        else:
            image = _image(data)
            visual_sources.append({"page": 1, "source_id": "page-1", "source": "image_page",
                                   "image": _jpeg_bytes(image)})
            if source_images is not None:
                source_images[1] = _jpeg_bytes(image)
                text, confidence = "", None
            else:
                text, confidence = _ocr_image(image)
            pages = [{"number": 1, "label": "Ảnh 1", "text": text,
                      "method": "ocr", "confidence": confidence, "image_source": "image_page"}]
        return _with_visuals(_finalize_pages(filename, pages, source_images), visual_sources)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("Không đọc được tệp. Kiểm tra định dạng và thử lại.") from exc


def extract_image_archive(data: bytes) -> dict:
    """Internal ordered archive created by JobQueue, never a public upload format."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if not 1 <= len(entries) <= 12 or any(entry.file_size > MAX_FILE_BYTES for entry in entries):
            raise ExtractionError("Bộ ảnh không hợp lệ")
        return extract_image_batch([(entry.filename, archive.read(entry)) for entry in entries])


def render_visual_source(filename: str, data: bytes, page_number: int,
                         source: dict, bbox: list[int] | None = None) -> bytes:
    """Render a stored source in display orientation, optionally crop 0..1000 box."""
    suffix = Path(filename).suffix.lower()
    kind = source.get("source")
    if kind == "pdf_page" and suffix == ".pdf":
        import pypdfium2 as pdfium

        document = pdfium.PdfDocument(data)
        try:
            if not 1 <= page_number <= len(document):
                raise ExtractionError("Không tìm thấy trang", 404)
            image = document[page_number - 1].render(scale=1.5).to_pil().convert("RGB")
        finally:
            document.close()
    elif kind == "office_page" and suffix in {".doc", ".docx", ".odt"}:
        _, page_images, _ = _office_page_images(data, suffix)
        if page_number not in page_images:
            raise ExtractionError("Could not find the requested rendered page.", 404)
        image = _image(page_images[page_number])
    elif kind == "image_page" and suffix in IMAGE_SUFFIXES and page_number == 1:
        image = _image(data)
    elif kind == "image_page" and suffix == ".imagebatch":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if not 1 <= page_number <= len(entries):
                raise ExtractionError("Không tìm thấy ảnh", 404)
            image = _image(archive.read(entries[page_number - 1]))
    elif kind == "docx_media" and suffix == ".docx":
        path = source.get("media_path", "")
        if not isinstance(path, str) or not path.startswith("word/media/") or ".." in path:
            raise ExtractionError("Đường dẫn hình không hợp lệ", 404)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            image = _image(archive.read(path))
    else:
        raise ExtractionError("Không tìm thấy hình", 404)
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise ExtractionError("Hình quá lớn để hiển thị", 422)
    if bbox is not None:
        if (len(bbox) != 4 or any(not isinstance(value, (int, float)) or not 0 <= value <= 1000
                                  for value in bbox) or bbox[2] <= bbox[0] or bbox[3] <= bbox[1]):
            raise ExtractionError("Vùng hình không hợp lệ", 422)
        left, top, right, bottom = (round(bbox[0] * image.width / 1000),
                                    round(bbox[1] * image.height / 1000),
                                    round(bbox[2] * image.width / 1000),
                                    round(bbox[3] * image.height / 1000))
        image = image.crop((left, top, right, bottom))
    return _jpeg_bytes(image)


def extract_text(filename: str, data: bytes) -> str:
    """Compatibility helper for old callers/tests."""
    return extract_document(filename, data)["text"]


if __name__ == "__main__":
    import json
    import sys
    from .config import positive_integer

    if sys.platform == "linux":
        import resource
        memory = positive_integer("WORKER_PARSER_MEMORY_MB", 512) * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (memory, memory))

    # Only the file extension is an argument; content travels over stdin and is
    # never logged. The parent worker enforces the process execution deadline.
    try:
        if sys.argv[1] == ".imagebatch":
            output = extract_image_archive(sys.stdin.buffer.read(12 * MAX_FILE_BYTES + 4097))
        else:
            output = extract_document("document" + sys.argv[1], sys.stdin.buffer.read(MAX_FILE_BYTES + 1))
    except ExtractionError as error:
        output = {"error": str(error), "status_code": error.status_code}
    sys.stdout.buffer.write(json.dumps(output, ensure_ascii=False).encode("utf-8"))
