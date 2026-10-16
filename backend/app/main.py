import hashlib
import io
import json
import os
import zipfile
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse

from .db import Database
from .schemas import AssignmentInput, RubricInput, SessionInput, ActionInput, MessageInput, SubmissionInput, ReviewInput
from .tutor import SEED, BLOOM_NAMES, lesson_steps, tutor_turn

ROOT = Path(__file__).resolve().parents[2]


def now():
    return datetime.now(timezone.utc).isoformat()


def require(conn, collection, identifier):
    value = Database.get(conn, collection, identifier)
    if value is None:
        raise HTTPException(404, "Không tìm thấy dữ liệu")
    return value


def public_session(session):
    return {key: value for key, value in session.items() if key != "requests"}


def saved_sessions(conn):
    return sorted(Database.all(conn, "sessions"), key=lambda s: s["updated_at"], reverse=True)


def create_app(database_path: str | None = None, frontend_dist: Path | None = None):
    db = Database(database_path or os.getenv("DATABASE_PATH", str(ROOT / "backend/data/socratic.db")))

    @asynccontextmanager
    async def lifespan(app):
        db.initialize()
        yield

    app = FastAPI(title="SocraVăn API", version="0.1.0", lifespan=lifespan,
                  description="Backend local/demo cho giao diện Socratic × Bloom. Gia sư dùng kịch bản, chưa kết nối LLM. Chưa có xác thực người dùng.")
    app.state.db = db
    app.add_middleware(CORSMiddleware, allow_origins=os.getenv(
        "CORS_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173"
    ).split(","), allow_methods=["GET", "POST", "PUT"], allow_headers=["Content-Type"])

    @app.get("/api/health", tags=["System"])
    def health():
        return {"status": "ok", "tutor_mode": "template", "auth_mode": "local_demo"}

    @app.get("/api/assignments", tags=["Assignments"])
    def assignments():
        with db.connect() as conn:
            items = db.all(conn, "assignments")
            sessions = saved_sessions(conn)
            for a in items:
                session = next((s for s in sessions if s["assignment_id"] == a["id"]), None)
                a["status"] = "Đã nộp" if session and session["submitted"] else "Đang học" if session else "Chưa bắt đầu"
                a["progress"] = 100 if session and session["submitted"] else round(session["step"] / len(session["steps"]) * 100) if session else 0
            return items

    @app.post("/api/assignments", status_code=201, tags=["Assignments"])
    def create_assignment(body: AssignmentInput):
        item = {**body.model_dump(), "id": str(uuid4()), "status": "Chưa bắt đầu", "progress": 0}
        with db.connect(write=True) as conn:
            db.put(conn, "assignments", item["id"], item)
        return item

    @app.post("/api/uploads/extract", tags=["Assignments"])
    def extract(file: UploadFile = File(...)):
        suffix = Path(file.filename or "").suffix.lower()
        if suffix not in {".txt", ".pdf", ".docx"}:
            raise HTTPException(415, "Chỉ hỗ trợ TXT, PDF có chữ và DOCX. Ảnh và PDF scan cần OCR, chưa được hỗ trợ.")
        data = file.file.read(10 * 1024 * 1024 + 1)
        if len(data) > 10 * 1024 * 1024:
            raise HTTPException(413, "Tệp vượt quá 10MB")
        try:
            if suffix == ".txt":
                text = data.decode("utf-8-sig")
            elif suffix == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(data))
                if reader.is_encrypted or len(reader.pages) > 100:
                    raise ValueError("PDF phải không mã hóa và tối đa 100 trang")
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
            else:
                from docx import Document
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                        raise ValueError("DOCX có kích thước giải nén quá lớn")
                doc = Document(io.BytesIO(data))
                text = "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
            text = text.strip()
            if len(text) < 15:
                raise ValueError("Không đủ văn bản để tạo đề. PDF scan cần OCR; hãy nhập nội dung trực tiếp.")
            if len(text) > 30000:
                raise ValueError("Nội dung quá dài (tối đa 30.000 ký tự); hãy chia nhỏ đề bài.")
        except Exception as exc:
            raise HTTPException(422, str(exc) if isinstance(exc, ValueError) else "Không đọc được tệp. Kiểm tra định dạng và thử lại.") from exc
        return {"filename": file.filename, "title": Path(file.filename or "Đề bài").stem, "text": text, "needs_confirmation": True}

    @app.get("/api/rubric", tags=["Rubric"])
    def rubric():
        with db.connect() as conn:
            return require(conn, "settings", "rubric")

    @app.put("/api/rubric", tags=["Rubric"])
    def save_rubric(body: RubricInput):
        with db.connect(write=True) as conn:
            current = require(conn, "settings", "rubric")
            current = {"version": current["version"] + 1, "criteria": [c.model_dump() for c in body.criteria]}
            db.put(conn, "settings", "rubric", current)
            return current

    @app.post("/api/sessions", tags=["Sessions"])
    def start_session(body: SessionInput):
        with db.connect(write=True) as conn:
            assignment = require(conn, "assignments", body.assignment_id)
            # Repeated starts (including a remount or double click) resume the saved session.
            existing = next((s for s in saved_sessions(conn) if s["assignment_id"] == body.assignment_id), None)
            if existing:
                return public_session(existing)
            steps = lesson_steps(assignment)
            session = {
                "id": str(uuid4()), "assignment_id": assignment["id"], "assignment": assignment,
                "student_id": "s1", "student_name": "Nguyễn Minh Anh", "class_name": "12A1",
                "steps": steps, "step": 0, "hint_tier": 0, "done": False, "submitted": False,
                "essay": "", "turns": [tutor_turn(0, steps)], "version": 0, "requests": {},
                "rubric": require(conn, "settings", "rubric"), "created_at": now(), "updated_at": now(),
                "tutor_mode": "template", "review": None,
            }
            db.put(conn, "sessions", session["id"], session)
            return public_session(session)

    @app.get("/api/sessions", tags=["Sessions"])
    def sessions():
        with db.connect() as conn:
            return [public_session(s) for s in saved_sessions(conn)]

    @app.get("/api/sessions/{session_id}", tags=["Sessions"])
    def get_session(session_id: str):
        with db.connect() as conn:
            return public_session(require(conn, "sessions", session_id))

    def mutate(session_id, body, action):
        with db.connect(write=True) as conn:
            session = require(conn, "sessions", session_id)
            fingerprint = hashlib.sha256(json.dumps({"action": action, **body.model_dump()}, sort_keys=True).encode()).hexdigest()
            previous = session["requests"].get(body.client_message_id)
            if previous:
                if previous != fingerprint:
                    raise HTTPException(409, "Mã yêu cầu đã được dùng cho nội dung khác")
                return public_session(session)
            if session["version"] != body.expected_version:
                raise HTTPException(409, "Phiên đã thay đổi. Vui lòng tải lại trước khi gửi tiếp.")
            if session["submitted"] or (session["done"] and action != "submit"):
                raise HTTPException(409, "Phiên đã kết thúc đối thoại")
            step = session["step"]
            current = session["steps"][step]
            if action == "hint":
                if session["hint_tier"] >= 3:
                    raise HTTPException(409, "Đã dùng hết gợi ý ở câu hỏi này")
                tier = session["hint_tier"] + 1
                session["turns"].append({"kind": "hint", "step": step, "tier": tier, "text": current["hints"][tier - 1]})
                session["hint_tier"] = tier
            elif action == "message":
                session["turns"].append({"kind": "student", "step": step, "text": body.text, "level": current["level"], "hints": session["hint_tier"], "quality": None})
                # Demo follows a question script. Progress is NOT evidence of mastery.
                if step < len(session["steps"]) - 1:
                    session["step"] += 1
                    session["hint_tier"] = 0
                    session["turns"].append(tutor_turn(session["step"], session["steps"]))
                else:
                    session["done"] = True
            else:
                if not session["done"]:
                    raise HTTPException(409, "Hoàn thành các câu hỏi trước khi nộp bài")
                session.update(essay=body.essay, submitted=True, submitted_at=now())
            session["version"] += 1
            session["updated_at"] = now()
            session["requests"][body.client_message_id] = fingerprint
            db.put(conn, "sessions", session_id, session)
            return public_session(session)

    @app.post("/api/sessions/{session_id}/messages", tags=["Sessions"])
    def message(session_id: str, body: MessageInput):
        return mutate(session_id, body, "message")

    @app.post("/api/sessions/{session_id}/hints", tags=["Sessions"])
    def hint(session_id: str, body: ActionInput):
        return mutate(session_id, body, "hint")

    @app.post("/api/sessions/{session_id}/submit", tags=["Sessions"])
    def submit(session_id: str, body: SubmissionInput):
        return mutate(session_id, body, "submit")

    @app.post("/api/sessions/{session_id}/reviews", tags=["Teacher"])
    def review(session_id: str, body: ReviewInput):
        with db.connect(write=True) as conn:
            session = require(conn, "sessions", session_id)
            if not session["submitted"]:
                raise HTTPException(409, "Bài làm chưa được nộp")
            criteria = session["rubric"]["criteria"]
            if set(body.scores) != {c["id"] for c in criteria}:
                raise HTTPException(422, "Cần chấm đầy đủ tiêu chí theo rubric của phiên học")
            result = {"id": str(uuid4()), "session_id": session_id, **body.model_dump(),
                      "total": round(sum(body.scores[c["id"]] * c["weight"] / 100 for c in criteria), 2),
                      "created_at": now(), "rubric_version": session["rubric"]["version"]}
            db.put(conn, "reviews", result["id"], result)
            session["review"] = result
            session["version"] += 1
            session["updated_at"] = now()
            db.put(conn, "sessions", session_id, session)
            return public_session(session)

    @app.get("/api/dashboard", tags=["Analytics"])
    def dashboard():
        with db.connect() as conn:
            saved = saved_sessions(conn)
            answers = [t for s in saved for t in s["turns"] if t["kind"] == "student"]
            hints = sum(t["kind"] == "hint" for s in saved for t in s["turns"])
            reviews = db.all(conn, "reviews")
            return {"sessions": len(saved), "answers": len(answers), "hints": hints,
                    "hints_per_answer": round(hints / len(answers), 1) if answers else 0,
                    "submitted": sum(s["submitted"] for s in saved),
                    "pending": sum(s["submitted"] and not s["review"] for s in saved),
                    "latest_session": public_session(next((s for s in saved if not s["submitted"]), saved[0])) if saved else None,
                    "latest_review": reviews[0] if reviews else None}

    @app.get("/api/progress", tags=["Analytics"])
    def progress():
        with db.connect() as conn:
            saved = saved_sessions(conn)
            reviewed = sorted([s for s in saved if s["review"]], key=lambda s: s["review"]["created_at"])
            latest = reviewed[-1] if reviewed else None
            return {
                "series": [{"week": f"Bài {i + 1}", "score": s["review"]["total"]} for i, s in enumerate(reviewed)],
                "bloom_profile": [{"skill": name, "value": sum(t["kind"] == "student" and t["level"] == i + 1 for s in saved for t in s["turns"])} for i, name in enumerate(BLOOM_NAMES)],
                "latest": public_session(latest) if latest else None,
                "pending": sum(s["submitted"] and not s["review"] for s in saved),
            }

    @app.get("/api/teacher/overview", tags=["Teacher"])
    def teacher_overview():
        with db.connect() as conn:
            saved = saved_sessions(conn)
            students = []
            for fixture in SEED["students"]:
                items = [s for s in saved if s["student_id"] == fixture["id"]]
                answers = [t for s in items for t in s["turns"] if t["kind"] == "student"]
                scores = [s["review"]["total"] for s in items if s["review"]]
                students.append({"id": fixture["id"], "name": fixture["name"], "cls": fixture["cls"],
                                 "sessions": len(items), "bloom": None, "trend": "—",
                                 "score": round(sum(scores) / len(scores), 1) if scores else None,
                                 "hints": round(sum(t["kind"] == "hint" for s in items for t in s["turns"]) / len(answers), 1) if answers else 0})
            return {"students": students, "sessions": len(saved), "pending": sum(s["submitted"] and not s["review"] for s in saved),
                    "distribution": [{"b": name, "n": sum(t["kind"] == "student" and t["level"] == i + 1 for s in saved for t in s["turns"])} for i, name in enumerate(BLOOM_NAMES)]}

    dist = (frontend_dist or ROOT / "frontend/dist").resolve()
    if dist.is_dir():
        @app.get("/{path:path}", include_in_schema=False)
        def frontend(path: str):
            if path == "api" or path.startswith("api/"):
                raise HTTPException(404, "API không tồn tại")
            candidate = (dist / path).resolve()
            if not candidate.is_relative_to(dist):
                raise HTTPException(404)
            if candidate.is_file():
                return FileResponse(candidate)
            if path.startswith("assets/") or Path(path).suffix:
                raise HTTPException(404)
            return FileResponse(dist / "index.html")

    return app


app = create_app()
