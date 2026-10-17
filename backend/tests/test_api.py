import io
from concurrent.futures import ThreadPoolExecutor
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from backend.app.main import create_app


@pytest.fixture
def client(database):
    with TestClient(create_app(database.path)) as c:
        yield c


def start(client, assignment="vo-nhat"):
    result = client.post("/api/sessions", json={"assignment_id": assignment})
    assert result.status_code == 200, result.text
    return result.json()


def action(client, session, endpoint, **payload):
    return client.post(f"/api/sessions/{session['id']}/{endpoint}", json={
        "client_message_id": str(uuid4()), "expected_version": session["version"], **payload,
    })


def complete(client, session):
    while not session["done"]:
        result = action(client, session, "messages", text="Em cần tìm dẫn chứng trong văn bản để giải thích ý kiến của mình.")
        assert result.status_code == 200, result.text
        session = result.json()
    result = action(client, session, "submit", essay="Bài làm độc lập của học sinh với luận điểm và dẫn chứng cụ thể để giáo viên đánh giá.")
    assert result.status_code == 200, result.text
    return result.json()


def test_full_student_teacher_flow(client):
    assert client.get("/api/health").json()["tutor_mode"] == "template"
    assignments = client.get("/api/assignments").json()
    assert len(assignments) == 3
    assert all(a["status"] == "Chưa bắt đầu" for a in assignments)
    assert client.get("/api/progress").json()["latest"] is None
    session = start(client)
    assert "requests" not in session
    assert start(client)["id"] == session["id"]
    hinted = action(client, session, "hints")
    assert hinted.status_code == 200
    session = complete(client, hinted.json())
    assert client.get("/api/dashboard").json()["pending"] == 1
    scores = {c["id"]: 8 for c in session["rubric"]["criteria"]}
    result = client.post(f"/api/sessions/{session['id']}/reviews", json={"scores": scores, "feedback": "Luận điểm rõ, cần giải thích kỹ hơn mối liên hệ giữa dẫn chứng và nhận định."})
    assert result.status_code == 200, result.text
    assert result.json()["review"]["total"] == 8
    progress = client.get("/api/progress").json()
    assert progress["latest"]["essay"] == session["essay"]
    assert progress["series"] == [{"week": "Bài 1", "score": 8}]
    assert sum(p["value"] for p in progress["bloom_profile"]) == 5
    assert client.get("/api/dashboard").json()["pending"] == 0
    overview = client.get("/api/teacher/overview").json()
    assert overview["students"][0]["score"] == 8
    assert overview["students"][0]["bloom"] is None
    assert all(t["quality"] is None for t in session["turns"] if t["kind"] == "student")


def test_idempotency_stale_versions_and_hint_limits(client):
    session = start(client)
    body = {"expected_version": 0, "client_message_id": "retry-me"}
    path = f"/api/sessions/{session['id']}/hints"
    first = client.post(path, json=body)
    repeated = client.post(path, json=body)
    assert first.json() == repeated.json()
    assert repeated.json()["hint_tier"] == 1
    assert action(client, session, "hints").status_code == 409
    assert client.post(f"/api/sessions/{session['id']}/messages", json={**body, "text": "Khác nội dung"}).status_code == 409
    session = first.json()
    for _ in range(2):
        session = action(client, session, "hints").json()
    assert action(client, session, "hints").status_code == 409
    assert client.get(path.removesuffix("/hints")).json()["hint_tier"] == 3


def test_concurrent_messages_do_not_overwrite_turns(client):
    session = start(client)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: action(client, session, "messages", text="Em chọn một chi tiết để giải thích."), range(2)))
    assert sorted(r.status_code for r in results) == [200, 409]
    saved = client.get(f"/api/sessions/{session['id']}").json()
    assert len([t for t in saved["turns"] if t["kind"] == "student"]) == 1


def test_persistence_across_restart(tmp_path):
    path = str(tmp_path / "persistent.db")
    with TestClient(create_app(path)) as first:
        session = start(first, "song")
        session = action(first, session, "messages", text="Em đang tìm hiểu yêu cầu đề bài Sóng.").json()
    with TestClient(create_app(path)) as second:
        assert second.get(f"/api/sessions/{session['id']}").json() == session
        assert len(second.get("/api/assignments").json()) == 3
        assert start(second, "song")["id"] == session["id"]


def test_assignment_selection_and_rubric_snapshot(client):
    created = client.post("/api/assignments", json={"title": "Trách nhiệm với cộng đồng", "prompt": "Hãy trình bày quan điểm của em về trách nhiệm với cộng đồng.", "target": 3})
    assert created.status_code == 201
    session = start(client, created.json()["id"])
    assert session["assignment"]["title"] == created.json()["title"]
    assert len(session["steps"]) == 3
    assert "Tràng" not in session["turns"][0]["text"]
    new_criteria = [{"id": "new", "name": "Lập luận", "weight": 100, "desc": "Luận điểm và dẫn chứng"}]
    result = client.put("/api/rubric", json={"criteria": new_criteria})
    assert result.status_code == 200
    assert result.json()["version"] == 2
    assert client.get(f"/api/sessions/{session['id']}").json()["rubric"]["version"] == 1
    assert start(client, "chi-pheo")["rubric"]["version"] == 2
    completed = complete(client, session)
    bad = client.post(f"/api/sessions/{completed['id']}/reviews", json={"scores": {"new": 9}, "feedback": "Cần dùng rubric gốc"})
    assert bad.status_code == 422


def test_validation_and_state_transitions(client):
    assert client.post("/api/assignments", json={"title": "X", "prompt": "Y"}).status_code == 422
    assert client.post("/api/assignments", json={"title": "Đề bài hợp lệ", "prompt": "Nội dung yêu cầu đủ dài.", "target": 7}).status_code == 422
    assert client.put("/api/rubric", json={"criteria": [{"id": "x", "name": "X", "weight": 90}]}).status_code == 422
    assert client.put("/api/rubric", json={"criteria": [{"id": "x", "name": "X", "weight": 50}] * 2}).status_code == 422
    assert client.get("/api/sessions/missing").status_code == 404
    assert client.post("/api/sessions", json={"assignment_id": "missing"}).status_code == 404
    session = start(client)
    assert action(client, session, "messages", text="   ").status_code == 422
    assert action(client, session, "submit", essay="Bài làm có đủ độ dài nhưng chưa hoàn tất câu hỏi.").status_code == 409
    assert client.post(f"/api/sessions/{session['id']}/reviews", json={"scores": {"c1": 8}, "feedback": "Chưa nộp bài"}).status_code == 409
    completed = complete(client, session)
    assert action(client, completed, "messages", text="Muốn sửa lịch sử").status_code == 409
    assert action(client, completed, "hints").status_code == 409
    scores = {c["id"]: 11 for c in completed["rubric"]["criteria"]}
    assert client.post(f"/api/sessions/{session['id']}/reviews", json={"scores": scores, "feedback": "Điểm vượt giới hạn"}).status_code == 422


def test_upload_extraction_requires_confirmation(client):
    result = client.post("/api/uploads/extract", files={"file": ("de-bai.txt", "Hãy phân tích vai trò của dẫn chứng trong bài nghị luận.".encode("utf-8"), "text/plain")})
    assert result.status_code == 200
    assert result.json()["needs_confirmation"] is True
    assert "dẫn chứng" in result.json()["text"]
    assert result.json()["structure"]["parts"][0]["questions"][0]["label"] == "Nội dung chưa tách được"
    assert len(client.get("/api/assignments").json()) == 3
    assert client.post("/api/uploads/extract", files={"file": ("scan.png", b"image", "image/png")}).status_code == 422
    assert client.post("/api/uploads/extract", files={"file": ("broken.pdf", b"not pdf", "application/pdf")}).status_code == 422
    assert client.post("/api/uploads/extract", files={"file": ("large.txt", b"x" * (10 * 1024 * 1024 + 1), "text/plain")}).status_code == 413


def test_structure_endpoint_keeps_only_selected_question_context(client):
    text = "PHẦN I. ĐỌC HIỂU\nVăn bản dùng chung.\nCâu 1. Tìm hình ảnh.\nCâu 2. Nêu tác dụng."
    result = client.post("/api/uploads/structure", json={"text": text})
    assert result.status_code == 200
    assert len(result.json()["parts"][0]["questions"]) == 2
    assert "Văn bản dùng chung" in result.json()["parts"][0]["context"]
    assert client.post("/api/uploads/structure", json={"text": "ngắn"}).status_code == 422


def test_docx_extraction(client):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Đề bài: Phân tích một vấn đề nghị luận xã hội.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Yêu cầu có dẫn chứng cụ thể."
    doc.add_paragraph("Câu 1. Nêu một dẫn chứng phù hợp.")
    stream = io.BytesIO()
    doc.save(stream)
    response = client.post("/api/uploads/extract", files={"file": ("de-bai.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200
    assert "dẫn chứng" in response.json()["text"]
    assert response.json()["text"].index("Yêu cầu có dẫn chứng") < response.json()["text"].index("Câu 1")
    assert response.json()["pages"][0]["label"] == "Tài liệu Word"
    assert response.json()["structure"]["parts"][0]["questions"][0]["label"] == "Câu 1"


def test_image_upload_returns_ocr_structure(client, monkeypatch):
    from PIL import Image
    from backend.app import extraction

    picture = Image.new("RGB", (480, 300), "white")
    stream = io.BytesIO()
    picture.save(stream, format="JPEG")
    monkeypatch.setattr(extraction, "_ocr_image", lambda _: ("PHẦN I. ĐỌC HIỂU\nCâu 1. Em thấy gì?", 90.0))
    response = client.post("/api/uploads/extract", files={"file": ("anh-1.jpg", stream.getvalue(), "image/jpeg")})
    assert response.status_code == 200
    assert response.json()["pages"][0]["method"] == "ocr"
    assert response.json()["structure"]["parts"][0]["questions"][0]["label"] == "Câu 1"


def test_static_frontend_and_unknown_api(tmp_path):
    dist = tmp_path / "dist"
    dist.mkdir()
    (dist / "index.html").write_text("<html>SocraVan</html>")
    (dist / "app.js").write_text("console.log('hello')")
    with TestClient(create_app(str(tmp_path / "test.db"), frontend_dist=dist)) as c:
        assert "SocraVan" in c.get("/phien-hoc").text
        assert c.get("/app.js").status_code == 200
        assert c.get("/api/nonexistent").status_code == 404
        assert c.get("/assets/missing.js").status_code == 404
