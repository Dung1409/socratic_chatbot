import io

from PIL import Image
from docx import Document
from fastapi.testclient import TestClient

from backend.app import extraction, visual_layout
from backend.app.main import create_app
from backend.app.structure import structure_text


def _png(width=200, height=100):
    stream = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(stream, format="PNG")
    return stream.getvalue()


def test_visual_regions_are_validated_before_being_returned():
    values = visual_layout._normalized({"regions": [
        {"kind": "graph", "bbox": [100, 200, 900, 700], "orientation": "horizontal",
         "question_label": "Câu 2", "description": "Trục thời gian và vận tốc", "uncertain": False},
        {"kind": "graph", "bbox": [900, 200, 100, 700]},
        {"kind": "answer", "bbox": [0, 0, 1000, 1000]},
    ]}, 2, "page-2", "pdf_page")
    assert len(values) == 1
    assert values[0]["id"] == "page-2-region-1"
    assert values[0]["bbox"] == [100, 200, 900, 700]
    assert values[0]["link_status"] == "unreviewed"


def test_question_links_are_only_proposed_for_unique_label_on_the_same_page():
    text = "I. TRẮC NGHIỆM\nCâu 1. Chọn đồ thị đúng.\nCâu 2. Chọn đồ thị sai."
    pages = [{"number": 1, "label": "Trang 1", "text": text,
              "visuals": [{"id": "page-1-region-1", "page": 1, "question_label": "Câu 2"},
                          {"id": "page-1-region-2", "page": 1, "question_label": None}]}]
    result = structure_text(text, pages)
    first, second = result["parts"][0]["questions"]
    assert result["visuals"][0]["question_ids"] == [second["id"]]
    assert first["visual_candidates"] == ["page-1-region-1", "page-1-region-2"]
    assert second["visual_refs"] == ["page-1-region-1"]
    assert result["visuals"][1]["question_ids"] == []


def test_figure_before_questions_remains_shared_candidate_not_forced_link():
    pages = [
        {"number": 1, "label": "Trang 1", "text": "I. TRẮC NGHIỆM\nQuan sát sơ đồ sau:",
         "visuals": [{"id": "page-1-region-1", "page": 1, "question_label": "Câu 1"}]},
        {"number": 2, "label": "Trang 2", "text": "Câu 1. Nhận xét sơ đồ.\nCâu 2. Chọn phương án đúng."},
    ]
    result = structure_text("\n\n".join(page["text"] for page in pages), pages)
    question_ids = [question["id"] for question in result["parts"][0]["questions"]]
    assert result["visuals"][0]["candidate_question_ids"] == question_ids
    assert result["visuals"][0]["question_ids"] == []


def test_docx_embedded_image_is_kept_as_addressable_source(monkeypatch):
    monkeypatch.setattr(visual_layout, "enabled", lambda: False)
    document = Document()
    document.add_paragraph("Câu 1. Quan sát hình sau và giải thích hiện tượng.")
    document.add_picture(io.BytesIO(_png()))
    stream = io.BytesIO()
    document.save(stream)
    data = stream.getvalue()
    extracted = extraction.extract_document("de.docx", data)
    source = extracted["pages"][0]["visual_sources"][0]
    assert source["source"] == "docx_media"
    assert extracted["pages"][0]["visuals"][0]["id"] == "media-1-region-1"
    image = extraction.render_visual_source("de.docx", data, 1, source)
    assert image.startswith(b"\xff\xd8")


def test_visual_crop_uses_oriented_image_coordinates():
    image = extraction.render_visual_source("de.png", _png(200, 100), 1,
                                            {"source": "image_page"}, [0, 0, 500, 1000])
    cropped = Image.open(io.BytesIO(image))
    assert cropped.size == (100, 100)


def test_phone_image_exif_rotation_is_applied_before_cropping():
    source = Image.new("RGB", (120, 60), "white")
    exif = source.getexif()
    exif[274] = 6  # Rotate 90 degrees clockwise for display.
    stream = io.BytesIO()
    source.save(stream, format="JPEG", exif=exif)
    image = extraction.render_visual_source("de.jpg", stream.getvalue(), 1,
                                            {"source": "image_page"}, [0, 0, 1000, 500])
    assert Image.open(io.BytesIO(image)).size == (60, 60)


def test_vision_failure_keeps_local_text_and_source(monkeypatch):
    monkeypatch.setattr(visual_layout, "enabled", lambda: True)
    monkeypatch.setattr(visual_layout, "analyze", lambda *args: (_ for _ in ()).throw(ValueError("bad JSON")))
    monkeypatch.setattr(extraction, "_ocr_image", lambda _: ("Câu 1. Hãy quan sát hình bên.", 90))
    result = extraction.extract_document("de.png", _png())
    assert result["pages"][0]["text"].startswith("Câu 1")
    assert result["pages"][0]["visual_sources"][0]["source_id"] == "page-1"
    assert result["pages"][0]["visuals"] == []


def test_uploaded_visual_is_private_and_can_be_cropped(database, tmp_path, authorize, monkeypatch):
    monkeypatch.setattr(visual_layout, "enabled", lambda: True)
    monkeypatch.setattr(visual_layout, "analyze", lambda *args: [
        {"id": "page-1-region-1", "page": 1, "source_id": "page-1", "source": "image_page",
         "kind": "graph", "bbox": [0, 0, 500, 1000], "question_label": "Câu 1",
         "description": "Đồ thị", "uncertain": False, "link_status": "unreviewed"}])
    monkeypatch.setattr(extraction, "_ocr_image", lambda _: ("Câu 1. Quan sát đồ thị và nêu nhận xét.", 90))
    storage = tmp_path / "visual-uploads"
    with TestClient(create_app(database.path, storage_path=storage)) as client:
        authorize(client)
        response = client.post("/api/uploads/extract", files={"file": ("de.png", _png(), "image/png")})
        assert response.status_code == 200, response.text
        document_id = response.json()["document_id"]
        visual = response.json()["structure"]["visuals"][0]
        assert visual["question_ids"] == ["part-1-question-1"]
        path = f"/api/documents/{document_id}/visuals/{visual['id']}/image"
        cropped = client.get(path)
        assert cropped.status_code == 200
        assert Image.open(io.BytesIO(cropped.content)).size == (100, 100)
        assert client.get(f"/api/documents/{document_id}/pages/1/image").status_code == 200
        with TestClient(create_app(database.path, storage_path=storage)) as other:
            authorize(other, email="another@example.com")
            assert other.get(path).status_code == 404
